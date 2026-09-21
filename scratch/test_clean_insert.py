import docx
from pathlib import Path
import win32com.client

template_path = r'C:\Users\HP\Desktop\MCA Project\docs\Template 1 MCA - Project I Documentation Template for Research Application Based Project.docx'
test_out = Path(r'c:\Users\HP\Desktop\MCA Project\MCA Project AI Breed\scratch\test_clean_doc.docx')
test_pdf = Path(r'c:\Users\HP\Desktop\MCA Project\MCA Project AI Breed\scratch\test_clean_doc.pdf')

doc = docx.Document(template_path)

# Let's inspect cleaning between Table 7 (Chapter 1) and Table 8 (Chapter 2)
t7 = doc.tables[7]._element
t8 = doc.tables[8]._element
parent = t7.getparent()

# Remove the empty placeholder paragraphs between t7 and t8
while True:
    next_sibling = t7.getnext()
    if next_sibling == t8:
        break
    parent.remove(next_sibling)

print("Placeholder paragraphs between T7 and T8 removed.")

# Now insert a page break after T7, then chapter content, then a page break before T8
# Create temporary paragraph in doc, then move it
p_br1 = doc.add_paragraph()
p_br1.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
t8.addprevious(p_br1._element)

p_head = doc.add_paragraph()
p_head.add_run("CHAPTER 1: INTRODUCTION").bold = True
t8.addprevious(p_head._element)

p_body = doc.add_paragraph()
p_body.add_run("This is the verified introduction to the AI-Powered Breed Recognition System.")
t8.addprevious(p_body._element)

p_br2 = doc.add_paragraph()
p_br2.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
t8.addprevious(p_br2._element)

doc.save(str(test_out))
print("Saved test_clean_doc.docx successfully!")

# Test Word COM conversion
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    doc_com = word.Documents.Open(str(test_out.resolve()))
    doc_com.SaveAs(str(test_pdf.resolve()), FileFormat=17)
    doc_com.Close()
    print(f"Word COM successfully compiled PDF: {test_pdf} (size: {test_pdf.stat().st_size} bytes)")
finally:
    word.Quit()
