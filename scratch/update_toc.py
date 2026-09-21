import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import win32com.client
from pathlib import Path

docx_path = Path(r'c:\Users\HP\Desktop\MCA Project\MCA Project AI Breed\MCA_Project_Documentation.docx')
pdf_path = Path(r'c:\Users\HP\Desktop\MCA Project\MCA Project AI Breed\MCA_Project_Documentation.pdf')

doc = docx.Document(docx_path)
t = doc.tables[7] # Table of Contents

def set_cell_font(p, text, size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.text = ""
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

# Part 1: Chapters 1 to 5
items_part1 = [
    ("1", "Chapter 1: Introduction", "19", True),
    ("", "  1.1 Introduction to the Project", "19", False),
    ("", "  1.2 Problem Statement", "20", False),
    ("", "  1.3 Objectives of the Study", "21", False),
    ("", "  1.4 Scope of the Project", "22", False),
    ("", "  1.5 Organization of the Report", "22", False),
    ("2", "Chapter 2: Literature Review", "25", True),
    ("", "  2.1 Review of Relevant Existing Approaches", "25", False),
    ("", "  2.2 Summary of Literature Review", "27", False),
    ("", "  2.3 Research Gap Analysis", "29", False),
    ("3", "Chapter 3: System Requirements", "32", True),
    ("", "  3.1 Software Requirements", "32", False),
    ("", "  3.2 Hardware Requirements", "32", False),
    ("", "  3.3 Tech Stack Used", "33", False),
    ("4", "Chapter 4: Proposed Methodology & System Design", "36", True),
    ("", "  4.1 Proposed Methodology", "36", False),
    ("", "  4.2 Modules and Description", "36", False),
    ("", "  4.3 System Architecture", "37", False),
    ("", "  4.4 Work Flow Diagram", "38", False),
    ("", "  4.5 Database Design", "39", False),
    ("", "  4.6 Input Design", "41", False),
    ("", "  4.7 Output Design", "42", False),
    ("5", "Chapter 5: System Implementation", "45", True),
    ("", "  5.1 Algorithm Implementation", "45", False),
    ("", "  5.2 Coding and Development", "46", False),
    ("", "  5.3 Implementation Tools", "46", False),
]

# Clear cell paragraphs
c0 = t.rows[1].cells[0]
c1 = t.rows[1].cells[1]
c2 = t.rows[1].cells[2]

c0.text = ""
c1.text = ""
c2.text = ""

for ch_no, content, page_no, is_bold in items_part1:
    p0 = c0.add_paragraph() if c0.paragraphs[0].text else c0.paragraphs[0]
    set_cell_font(p0, ch_no, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p1 = c1.add_paragraph() if c1.paragraphs[0].text else c1.paragraphs[0]
    set_cell_font(p1, content, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    p2 = c2.add_paragraph() if c2.paragraphs[0].text else c2.paragraphs[0]
    set_cell_font(p2, page_no, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.CENTER)

# Part 2: Chapters 6, 7, Bibliography, Appendices, Research Paper
items_part2 = [
    ("6", "Chapter 6: Results and Discussions", "49", True),
    ("", "  6.1 Performance Metrics Used", "49", False),
    ("", "  6.2 Implementation Results (Tables and Graphs)", "49", False),
    ("", "  6.3 Discussions", "54", False),
    ("7", "Chapter 7: Conclusion and Future Enhancement", "58", True),
    ("", "  7.1 Conclusion", "58", False),
    ("", "  7.2 Scope for Future Enhancement", "58", False),
    ("", "Bibliography / References", "60", True),
    ("", "Appendix", "62", True),
    ("", "  A. Key Source Code Listings", "62", False),
    ("", "  B. Output Screens and Visual Interfaces", "63", False),
    ("", "Research Paper Publication Status", "66", True),
]

c0_2 = t.rows[3].cells[0]
c1_2 = t.rows[3].cells[1]
c2_2 = t.rows[3].cells[2]

c0_2.text = ""
c1_2.text = ""
c2_2.text = ""

for ch_no, content, page_no, is_bold in items_part2:
    p0 = c0_2.add_paragraph() if c0_2.paragraphs[0].text else c0_2.paragraphs[0]
    set_cell_font(p0, ch_no, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p1 = c1_2.add_paragraph() if c1_2.paragraphs[0].text else c1_2.paragraphs[0]
    set_cell_font(p1, content, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    p2 = c2_2.add_paragraph() if c2_2.paragraphs[0].text else c2_2.paragraphs[0]
    set_cell_font(p2, page_no, size=10.5, bold=is_bold, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(str(docx_path))
print("Updated Table of Contents table with verified page numbers!")

# Recompile PDF via Word COM
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    doc_com = word.Documents.Open(str(docx_path.resolve()))
    doc_com.SaveAs(str(pdf_path.resolve()), FileFormat=17)
    doc_com.Close()
    print(f"Recompiled PDF with updated Table of Contents: {pdf_path} (Size: {pdf_path.stat().st_size} bytes)")
finally:
    word.Quit()
