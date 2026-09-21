"""Builder script to generate the Complete MCA Major Project Documentation.

Follows the official PSG College of Arts & Science MCA documentation template:
'Template 1 MCA - Project I Documentation Template for Research Application Based Project.docx'
Adheres strictly to 'General Instructions for Documentation.pdf'.
"""

import os
import shutil
from pathlib import Path
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

WORKSPACE = Path(r"c:\Users\HP\Desktop\MCA Project\MCA Project AI Breed")
TEMPLATE_PATH = Path(r"C:\Users\HP\Desktop\MCA Project\docs\Template 1 MCA - Project I Documentation Template for Research Application Based Project.docx")
OUTPUT_DOCX = WORKSPACE / "MCA_Project_Documentation.docx"
OUTPUT_PDF = WORKSPACE / "MCA_Project_Documentation.pdf"
ASSETS_DIR = WORKSPACE / "documentation_assets"

# Verified Project Constants
PROJECT_TITLE = "AI-POWERED INTELLIGENT BREED RECOGNITION SYSTEM FOR INDIAN CATTLE AND BUFFALOES USING DEEP LEARNING"
SUBTITLE_ALT = "DEEP LEARNING BASED CATTLE AND BUFFALO IMAGE CLASSIFICATION SYSTEM"
COLLEGE_NAME = "PSG COLLEGE OF ARTS & SCIENCE"
COLLEGE_AFFILIATION = "An Autonomous College - Affiliated to Bharathiar University"
COLLEGE_ACCREDITATION = "Accredited with 'A++' grade by NAAC (4th Cycle)"
DEPARTMENT_NAME = "DEPARTMENT OF COMPUTER APPLICATIONS (PG)"
DEGREE_NAME = "MASTER OF COMPUTER APPLICATIONS"
MONTH_YEAR = "OCTOBER 2026"

STUDENT_NAME_PLACEHOLDER = "[STUDENT NAME]"
REGISTER_NO_PLACEHOLDER = "[25MCA0XX]"
GUIDE_NAME_PLACEHOLDER = "[FACULTY GUIDE NAME]"
GUIDE_DEGREE_PLACEHOLDER = "[Degree, Designation]"
VIVA_DATE_PLACEHOLDER = "[Date of Viva-Voce Examination]"


def set_run_font(run, name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0, 0, 0)):
    """Set font styling on a docx run."""
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)


def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5):
    """Format paragraph alignment, spacing and line height."""
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing


def add_body_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, bold=False, italic=False):
    """Add standard justified 12pt Times New Roman body paragraph with 1.5 line spacing."""
    p = doc.add_paragraph()
    format_paragraph(p, align=align, space_before=space_before, space_after=space_after, line_spacing=1.5)
    run = p.add_run(text)
    set_run_font(run, name="Times New Roman", size_pt=12, bold=bold, italic=italic)
    return p


def add_heading_1(doc, text):
    """Add Chapter / Level 1 Heading (Left-aligned, 14pt, Bold)."""
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6, line_spacing=1.5)
    run = p.add_run(text)
    set_run_font(run, name="Times New Roman", size_pt=14, bold=True)
    return p


def add_heading_2(doc, text):
    """Add Section / Level 2 Heading (Left-aligned, 12pt, Bold)."""
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, line_spacing=1.5)
    run = p.add_run(text)
    set_run_font(run, name="Times New Roman", size_pt=12, bold=True)
    return p


def add_heading_3(doc, text):
    """Add Sub-section / Level 3 Heading (Left-aligned, 12pt, Bold/Italic)."""
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4, line_spacing=1.5)
    run = p.add_run(text)
    set_run_font(run, name="Times New Roman", size_pt=12, bold=True, italic=True)
    return p


def add_bullet_item(doc, bold_prefix, text):
    """Add bullet paragraph with bold prefix."""
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=2, space_after=4, line_spacing=1.5)
    r1 = p.add_run(bold_prefix)
    set_run_font(r1, name="Times New Roman", size_pt=12, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, name="Times New Roman", size_pt=12, bold=False)
    return p


def set_cell_background(cell, fill_hex):
    """Set shading color on table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Set subtle uniform borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_styled_table(doc, headers, data, col_widths=None, alignment=WD_TABLE_ALIGNMENT.CENTER):
    """Add a beautifully styled academic table with headers, borders and clean formatting."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = alignment
    set_table_borders(table, color="B0C4DE", sz="6")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=4, line_spacing=1.15)
        run = p.add_run(h)
        set_run_font(run, name="Times New Roman", size_pt=11, bold=True, color_rgb=(25, 25, 112))
        set_cell_background(hdr_cells[i], "E6F0FA")
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_fill = "FAFAFA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            # Align first col left, others center or justify
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            format_paragraph(p, align=align, space_before=3, space_after=3, line_spacing=1.15)
            run = p.add_run(str(val))
            set_run_font(run, name="Times New Roman", size_pt=10.5, bold=False)
            set_cell_background(row_cells[c_idx], bg_fill)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    format_paragraph(p_after, space_before=0, space_after=6, line_spacing=1.0)
    return table


def add_figure_image(doc, img_path, caption_text, width_inches=5.8):
    """Add centered figure with proper academic caption."""
    p_img = doc.add_paragraph()
    format_paragraph(p_img, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=4, line_spacing=1.0)
    if Path(img_path).exists():
        p_img.add_run().add_picture(str(img_path), width=Inches(width_inches))
    else:
        run_placeholder = p_img.add_run(f"[INSERT ACTUAL FIGURE/SCREENSHOT HERE: {caption_text}]")
        set_run_font(run_placeholder, name="Times New Roman", size_pt=11, bold=True, italic=True, color_rgb=(178, 34, 34))
    
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=10, line_spacing=1.15)
    run_cap = p_cap.add_run(caption_text)
    set_run_font(run_cap, name="Times New Roman", size_pt=10.5, bold=True, italic=True, color_rgb=(47, 79, 79))
    return p_img


def add_table_caption(doc, caption_text):
    """Add standard academic table caption above table."""
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4, line_spacing=1.15)
    run_cap = p_cap.add_run(caption_text)
    set_run_font(run_cap, name="Times New Roman", size_pt=11, bold=True, color_rgb=(25, 25, 112))
    return p_cap


def add_page_break(doc):
    """Add standard page break."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=0, line_spacing=1.0)
    p.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    return p


print("Helper functions defined successfully.")
